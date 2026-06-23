from __future__ import annotations

import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import BinaryIO

import pandas as pd
from openpyxl.styles import PatternFill


ENCODINGS = ("utf-8-sig", "utf-8", "latin1", "cp1252")
TOLERANCE = Decimal("0.01")

STATUS_OK = "OK"
STATUS_DIVERGENT = "DIVERGENTE"
STATUS_NOT_FOUND = "N\u00c3O ENCONTRADO NO TCE"
STATUS_EXTRA_TCE = "SOBRA NO TCE"

COL_CODIGO_BETHA = "C\u00f3digo Betha"
COL_EVENTO_BETHA = "Evento Betha"
COL_VALOR_BETHA = "Valor Betha"
COL_TIPO_BETHA = "Tipo"
COL_VALOR_TCE = "Valor TCE"
COL_DIFERENCA = "Diferen\u00e7a"
COL_STATUS = "Status"
COL_CODIGO_TCE_META = "__Codigo TCE"

RESULT_COLUMNS = [
    COL_CODIGO_BETHA,
    COL_EVENTO_BETHA,
    COL_VALOR_BETHA,
    COL_TIPO_BETHA,
    COL_VALOR_TCE,
    COL_DIFERENCA,
    COL_STATUS,
]


class ConferenciaError(Exception):
    """Erro esperado que pode ser exibido ao usuario pelo sistema chamador."""


def repair_broken_text(value: str) -> str:
    replacements = {
        "AUX?LIO": "AUXILIO",
        "DOEN?A": "DOENCA",
        "LICEN?A": "LICENCA",
        "DIFEREN?A": "DIFERENCA",
        "QUINQU?NIO": "QUINQUENIO",
        "F?RIAS": "FERIAS",
        "SUBS?DIO": "SUBSIDIO",
        "PECUNI?RIO": "PECUNIARIO",
        "SAL?RIO": "SALARIO",
        "C?DIGO": "CODIGO",
        "FUNCION?RIOS": "FUNCIONARIOS",
        "RESCIS?O": "RESCISAO",
        "UNI?O": "UNIAO",
    }
    text = str(value)
    for old, new in replacements.items():
        text = re.sub(re.escape(old), new, text, flags=re.IGNORECASE)
    return text


def normalize_text(value: object, *, strip_initial_code: bool = False) -> str:
    if pd.isna(value):
        return ""

    text = str(value).strip()
    if strip_initial_code:
        text = re.sub(r"^\s*\d+\s*[-\u2013.]\s*", "", text)

    text = repair_broken_text(text)
    text = unicodedata.normalize("NFKD", text)
    text = text.encode("ascii", "ignore").decode("ascii")
    text = text.upper()
    text = text.replace("?", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text.rstrip(" .")


def normalize_column_name(value: object) -> str:
    return normalize_text(value).replace(" ", "")


def extract_initial_code(value: object) -> str:
    match = re.match(r"^\s*(\d+)\s*[-\u2013.]?\s*", str(value or ""))
    return match.group(1) if match else ""


def split_codes(value: object) -> set[str]:
    return set(re.findall(r"\d+", str(value or "")))


def extract_value_type(value: object) -> str:
    match = re.search(r"\s*([PD])\s*$", str(value or ""), flags=re.IGNORECASE)
    return match.group(1).upper() if match else ""


def format_value_type(value: object) -> str:
    value_type = str(value or "").strip().upper()
    labels = {"P": "P - Provento", "D": "D - Desconto"}
    return labels.get(value_type, value_type)


def aggregate_value_types(values: pd.Series) -> str:
    value_types = sorted({extract_value_type(value) for value in values if extract_value_type(value)})
    return ", ".join(value_types)


def find_column(df: pd.DataFrame, candidates: set[str], label: str) -> str:
    normalized_columns = {normalize_column_name(column): column for column in df.columns}
    for candidate in candidates:
        if candidate in normalized_columns:
            return normalized_columns[candidate]

    available = ", ".join(str(column) for column in df.columns)
    raise ConferenciaError(f"Coluna obrigatoria nao encontrada: {label}. Colunas lidas: {available}")


def read_binary(source: str | Path | bytes | bytearray | BinaryIO) -> bytes:
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if hasattr(source, "getvalue"):
        return source.getvalue()
    if hasattr(source, "read"):
        return source.read()
    raise TypeError("Fonte de arquivo invalida. Use caminho, bytes ou arquivo binario.")


def decode_file(source: str | Path | bytes | bytearray | BinaryIO) -> tuple[str, str]:
    raw = read_binary(source)
    last_error: Exception | None = None

    for encoding in ENCODINGS:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError as exc:
            last_error = exc

    raise ConferenciaError(f"Nao foi possivel ler o arquivo com os encodings suportados: {last_error}")


def detect_header_line(lines: list[str]) -> int:
    for index, line in enumerate(lines[:80]):
        normalized = normalize_text(line)
        if "NOME COMPONENTE" in normalized and "VALOR" in normalized:
            return index
        if "EVENTO" in normalized and "VALOR CALCULADO" in normalized:
            return index
    return 0


def detect_delimiter(sample: str) -> str:
    first_lines = [line for line in sample.splitlines()[:10] if line.strip()]
    semicolon_count = sum(line.count(";") for line in first_lines)
    comma_count = sum(line.count(",") for line in first_lines)
    return ";" if semicolon_count >= comma_count else ","


def read_csv_auto(source: str | Path | bytes | bytearray | BinaryIO) -> tuple[pd.DataFrame, dict[str, object]]:
    text, encoding = decode_file(source)
    lines = text.splitlines()
    if not lines:
        raise ConferenciaError("Arquivo vazio.")

    header_line = detect_header_line(lines)
    sample = "\n".join(lines[header_line : header_line + 12])
    delimiter = detect_delimiter(sample)

    try:
        df = pd.read_csv(
            io.StringIO(text),
            sep=delimiter,
            skiprows=header_line,
            dtype=str,
            engine="python",
        )
    except Exception as exc:
        raise ConferenciaError(f"CSV invalido ou nao suportado: {exc}") from exc

    df = df.dropna(axis=1, how="all")
    df.columns = [str(column).strip() for column in df.columns]
    metadata = {"encoding": encoding, "delimiter": delimiter, "headerLine": header_line + 1}
    return df, metadata


def parse_brazilian_money(value: object, *, column: str) -> Decimal:
    if pd.isna(value):
        return Decimal("0")

    original = str(value).strip()
    if not original:
        return Decimal("0")

    cleaned = original.replace("R$", "").strip()
    cleaned = re.sub(r"\s*[PD]\s*$", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = cleaned.replace(" ", "")

    if not re.fullmatch(r"-?(\d{1,3}(\.\d{3})*|\d+)(,\d+)?|-?\d+(\.\d+)?", cleaned):
        raise ConferenciaError(f"Valor monetario invalido na coluna {column}: {original}")

    if "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")

    try:
        return Decimal(cleaned).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise ConferenciaError(f"Valor monetario invalido na coluna {column}: {original}") from exc


def prepare_betha(df: pd.DataFrame) -> pd.DataFrame:
    code_col = find_column(df, {"CODIGO", "CIGO", "COD"}, "Codigo Betha")
    event_col = find_column(df, {"EVENTO"}, "Evento")
    value_col = find_column(df, {"VALORCALCULADO"}, "Valor calculado")

    work = df[[code_col, event_col, value_col]].copy()
    work.columns = ["codigo", "evento", "valor"]
    work["evento"] = work["evento"].fillna("").astype(str).str.strip()
    work = work[work["evento"] != ""]
    work["chave"] = work["evento"].map(normalize_text)
    work["valor_betha"] = work["valor"].map(lambda item: parse_brazilian_money(item, column="Valor calculado"))

    return (
        work.groupby("chave", as_index=False)
        .agg(
            codigo_betha=("codigo", lambda values: ", ".join(str(value).strip() for value in values if str(value).strip())),
            evento_betha=("evento", "first"),
            valor_betha=("valor_betha", "sum"),
            tipo_betha=("valor", aggregate_value_types),
        )
        .reset_index(drop=True)
    )


def prepare_tce(df: pd.DataFrame) -> pd.DataFrame:
    component_col = find_column(df, {"NOMECOMPONENTE"}, "Nome Componente")
    value_col = find_column(df, {"VALOR"}, "Valor")

    work = df[[component_col, value_col]].copy()
    work.columns = ["componente", "valor"]
    work["componente"] = work["componente"].fillna("").astype(str).str.strip()
    work = work[work["componente"] != ""]
    work["codigo_tce"] = work["componente"].map(extract_initial_code)
    work["chave"] = work["componente"].map(lambda item: normalize_text(item, strip_initial_code=True))
    work = work[work["chave"] != ""]
    work["valor_tce"] = work["valor"].map(lambda item: parse_brazilian_money(item, column="Valor"))

    return (
        work.groupby("chave", as_index=False)
        .agg(
            componente_tce=("componente", "first"),
            codigo_tce=("codigo_tce", lambda values: ", ".join(sorted(set(str(value).strip() for value in values if str(value).strip())))),
            valor_tce=("valor_tce", "sum"),
        )
        .reset_index(drop=True)
    )


def comparar_dataframes(betha_df: pd.DataFrame, tce_df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, object]]:
    betha = prepare_betha(betha_df)
    tce = prepare_tce(tce_df)

    merged = betha.merge(tce, on="chave", how="outer")
    merged["valor_betha"] = merged["valor_betha"].fillna(Decimal("0"))
    merged["valor_tce"] = merged["valor_tce"].fillna(Decimal("0"))
    merged["diferenca"] = merged["valor_betha"] - merged["valor_tce"]

    def status(row: pd.Series) -> str:
        has_betha = pd.notna(row.get("evento_betha"))
        has_tce = pd.notna(row.get("componente_tce"))
        if has_betha and not has_tce:
            return STATUS_NOT_FOUND
        if has_tce and not has_betha:
            return STATUS_EXTRA_TCE
        if abs(row["diferenca"]) <= TOLERANCE:
            return STATUS_OK
        return STATUS_DIVERGENT

    merged[COL_STATUS] = merged.apply(status, axis=1)
    merged[COL_CODIGO_BETHA] = merged["codigo_betha"].fillna("")
    merged[COL_EVENTO_BETHA] = merged["evento_betha"].fillna(merged["componente_tce"].fillna(""))
    merged[COL_VALOR_BETHA] = merged["valor_betha"]
    merged[COL_TIPO_BETHA] = merged["tipo_betha"].fillna("").map(format_value_type)
    merged[COL_VALOR_TCE] = merged["valor_tce"]
    merged[COL_DIFERENCA] = merged["diferenca"]
    merged[COL_CODIGO_TCE_META] = merged["codigo_tce"].fillna("")

    result = merged[RESULT_COLUMNS + [COL_CODIGO_TCE_META]].sort_values([COL_STATUS, COL_EVENTO_BETHA], kind="stable")
    summary = {
        "totalBetha": sum(result[COL_VALOR_BETHA], Decimal("0")),
        "totalTce": sum(result[COL_VALOR_TCE], Decimal("0")),
        "diferencaTotal": sum(result[COL_DIFERENCA], Decimal("0")),
        "ok": int((result[COL_STATUS] == STATUS_OK).sum()),
        "divergente": int((result[COL_STATUS] == STATUS_DIVERGENT).sum()),
        "naoEncontrado": int((result[COL_STATUS] == STATUS_NOT_FOUND).sum()),
        "sobraTce": int((result[COL_STATUS] == STATUS_EXTRA_TCE).sum()),
    }
    return result.reset_index(drop=True), summary


def decimal_to_float(value: object) -> float:
    if isinstance(value, Decimal):
        return float(value)
    return float(value or 0)


def dataframe_to_records(df: pd.DataFrame) -> list[dict[str, object]]:
    records = []
    for row in df.to_dict(orient="records"):
        converted = {}
        for key, value in row.items():
            converted[key] = decimal_to_float(value) if isinstance(value, Decimal) else value
        records.append(converted)
    return records


def summary_to_json(summary: dict[str, object]) -> dict[str, object]:
    return {
        key: decimal_to_float(value) if isinstance(value, Decimal) else value
        for key, value in summary.items()
    }


def processar_conferencia(
    betha_source: str | Path | bytes | bytearray | BinaryIO,
    tce_source: str | Path | bytes | bytearray | BinaryIO,
) -> dict[str, object]:
    """Processa os dois CSVs e retorna um dicionario pronto para API/frontend."""

    betha_df, betha_metadata = read_csv_auto(betha_source)
    tce_df, tce_metadata = read_csv_auto(tce_source)
    result, summary = comparar_dataframes(betha_df, tce_df)

    return {
        "rows": dataframe_to_records(result),
        "summary": summary_to_json(summary),
        "metadata": {
            "betha": betha_metadata | {"rows": len(betha_df)},
            "tce": tce_metadata | {"rows": len(tce_df)},
        },
    }


def gerar_excel_conferencia(rows: list[dict[str, object]] | pd.DataFrame, summary: dict[str, object]) -> bytes:
    """Gera o XLSX da conferencia e retorna bytes para download no sistema chamador."""

    result = rows.copy() if isinstance(rows, pd.DataFrame) else pd.DataFrame(rows)
    for column in RESULT_COLUMNS:
        if column not in result.columns:
            result[column] = ""
    result = result[RESULT_COLUMNS]
    return build_excel(result, summary)


def gerar_excel_conferencia_destacado(rows: list[dict[str, object]] | pd.DataFrame, summary: dict[str, object]) -> bytes:
    """Gera XLSX com destaque visual por status e direcao da diferenca."""

    result = rows.copy() if isinstance(rows, pd.DataFrame) else pd.DataFrame(rows)
    for column in RESULT_COLUMNS + [COL_CODIGO_TCE_META]:
        if column not in result.columns:
            result[column] = ""
    return build_excel(result[RESULT_COLUMNS + [COL_CODIGO_TCE_META]], summary, highlighted=True)


def has_code_mismatch(row: pd.Series) -> bool:
    betha_codes = split_codes(row.get(COL_CODIGO_BETHA, ""))
    tce_codes = split_codes(row.get(COL_CODIGO_TCE_META, ""))
    return bool(betha_codes and tce_codes and betha_codes.isdisjoint(tce_codes))


def row_highlight_fill(row: pd.Series) -> PatternFill | None:
    status = normalize_text(row.get(COL_STATUS, ""))
    valor_betha = Decimal(str(row.get(COL_VALOR_BETHA, 0) or 0))
    valor_tce = Decimal(str(row.get(COL_VALOR_TCE, 0) or 0))

    fills = {
        "green": PatternFill("solid", fgColor="C6EFCE"),
        "blue": PatternFill("solid", fgColor="BDD7EE"),
        "red": PatternFill("solid", fgColor="FFC7CE"),
        "yellow": PatternFill("solid", fgColor="FFEB9C"),
        "orange": PatternFill("solid", fgColor="F4B183"),
        "purple": PatternFill("solid", fgColor="D9E1F2"),
    }

    if status == normalize_text(STATUS_NOT_FOUND):
        return fills["yellow"]
    if has_code_mismatch(row):
        return fills["orange"]
    if status == normalize_text(STATUS_OK):
        return fills["green"]
    if status == normalize_text(STATUS_DIVERGENT):
        return fills["blue"] if valor_tce > valor_betha else fills["red"]
    return None


def value_type_fill(value: object) -> PatternFill | None:
    value_type = str(value or "").strip().upper()
    if value_type.startswith("P"):
        return PatternFill("solid", fgColor="E2F0D9")
    if value_type.startswith("D"):
        return PatternFill("solid", fgColor="FCE4D6")
    return None


def build_excel(result: pd.DataFrame, summary: dict[str, object], highlighted: bool = False) -> bytes:
    output = io.BytesIO()
    excel_result = result.copy()
    excel_output = excel_result[RESULT_COLUMNS].copy()
    for column in (COL_VALOR_BETHA, COL_VALOR_TCE, COL_DIFERENCA):
        excel_output[column] = excel_output[column].map(decimal_to_float)

    resumo = pd.DataFrame(
        [
            {"Indicador": "Total Betha", "Valor": decimal_to_float(summary.get("totalBetha", 0))},
            {"Indicador": "Total TCE", "Valor": decimal_to_float(summary.get("totalTce", 0))},
            {"Indicador": "Diferenca total", "Valor": decimal_to_float(summary.get("diferencaTotal", 0))},
            {"Indicador": "Eventos OK", "Valor": int(summary.get("ok", 0))},
            {"Indicador": "Divergencias", "Valor": int(summary.get("divergente", 0))},
            {"Indicador": "Nao encontrados no TCE", "Valor": int(summary.get("naoEncontrado", 0))},
            {"Indicador": "Sobras no TCE", "Valor": int(summary.get("sobraTce", 0))},
        ]
    )
    sobras = excel_output[excel_output[COL_STATUS] == STATUS_EXTRA_TCE].copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        excel_output.to_excel(writer, sheet_name="Conferencia", index=False)
        resumo.to_excel(writer, sheet_name="Resumo", index=False)
        sobras.to_excel(writer, sheet_name="Sobras TCE", index=False)

        workbook = writer.book
        money_format = '"R$" #,##0.00'
        for sheet_name in ("Conferencia", "Sobras TCE"):
            sheet = workbook[sheet_name]
            header_columns = {cell.value: cell.column for cell in sheet[1]}
            for column in (COL_VALOR_BETHA, COL_VALOR_TCE, COL_DIFERENCA):
                column_index = header_columns.get(column)
                if not column_index:
                    continue
                for column_cells in sheet.iter_cols(min_col=column_index, max_col=column_index, min_row=2):
                    for cell in column_cells:
                        cell.number_format = money_format
            for column_cells in sheet.columns:
                width = max(len(str(cell.value or "")) for cell in column_cells) + 2
                sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 60)

        if highlighted:
            sheet = workbook["Conferencia"]
            header_columns = {cell.value: cell.column for cell in sheet[1]}
            type_column_indexes = [
                header_columns[COL_TIPO_BETHA]
            ] if COL_TIPO_BETHA in header_columns else []
            for row_index, (_, result_row) in enumerate(excel_result.iterrows(), start=2):
                fill = row_highlight_fill(result_row)
                if fill:
                    for cell in sheet[row_index]:
                        cell.fill = fill
                for column_index in type_column_indexes:
                    cell = sheet.cell(row=row_index, column=column_index)
                    fill = value_type_fill(cell.value)
                    if fill:
                        cell.fill = fill

        resumo_sheet = workbook["Resumo"]
        for cell in resumo_sheet["B"]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = money_format if cell.row <= 4 else "0"
        for column_cells in resumo_sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            resumo_sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 45)

    return output.getvalue()
